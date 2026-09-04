from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
SOURCE = BASE_DIR / "ViPA-Fi_report_vi.md"
OUTPUT = BASE_DIR / "ViPA-Fi_report_vi.docx"


def add_code_block(document, lines):
    text = "\n".join(lines).strip()
    if not text:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_text_with_bold(paragraph, text):
    while "**" in text:
        before, _, rest = text.partition("**")
        if before:
            paragraph.add_run(before)
        bold_text, _, text = rest.partition("**")
        run = paragraph.add_run(bold_text)
        run.bold = True
    if text:
        paragraph.add_run(text)


def add_markdown_line(document, line):
    if line.startswith("- "):
        paragraph = document.add_paragraph(style="List Bullet")
        add_text_with_bold(paragraph, line[2:])
        return

    if len(line) > 3 and line[0].isdigit() and ". " in line[:4]:
        paragraph = document.add_paragraph(style="List Number")
        add_text_with_bold(paragraph, line.split(". ", 1)[1])
        return

    paragraph = document.add_paragraph()
    add_text_with_bold(paragraph, line)


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    for style_name, size in [("Normal", 11), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)

    in_code = False
    code_lines = []

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if line.startswith("# "):
            title = document.add_heading(line[2:], level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        else:
            add_markdown_line(document, line)

    if code_lines:
        add_code_block(document, code_lines)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
